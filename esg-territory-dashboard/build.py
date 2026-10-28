"""
ESG Territory Hours - one-file builder
=====================================
Reads one Excel workbook and writes two offline HTML pages:
  ESG_Hours_Dashboard.html   full dashboard (territories, growth, rates, US tagging scenario planner)
  ESG_US_Story_v2.html       partner story (arrow keys move between sections)

Run:   py build.py                      (or double-click build.py)
       py build.py --excel "D:\\Data\\Hours.xlsx"

Needs Python 3.10+ with:  py -m pip install pandas openpyxl
Files:  build.py (this) + template.html (page design) + your workbook. Missing input tabs are added
to the workbook on the first run (a backup copy is kept).
"""
# ================================ SETTINGS ================================
# Put private values (real workbook path, your name/email) in local_settings.py
# with the same names - that file is ignored by Git.
EXCEL_FILE = "ESG_Hours_Sample.xlsx"   # file name next to build.py, or a full path

CONTACT_NAME = ""                      # sign-off at the end of both pages
CONTACT_EMAIL = ""
CONTACT_NAME_2 = ""                    # optional second contact
CONTACT_EMAIL_2 = ""
CONTACT_TEAM = ""                      # optional, shown after the contacts

LOGO_FILE = "logo.png"                 # optional logo next to build.py (PNG/JPG/SVG)
LOGO_TEXT = ""                         # text shown when there is no logo file
REPORT_TITLE = "ESG Delivery: Territory Hours"      # if the workbook has no 'Report title' assumption
CLASSIFICATION = "Confidential | Internal use only" # if the workbook has no 'Classification' assumption

TEMPLATE_FILE = "template.html"
DASHBOARD_FILE = "ESG_Hours_Dashboard.html"
STORY_FILE = "ESG_US_Story_v2.html"
OPEN_WHEN_DONE = True                  # open the story page at the end
# ==========================================================================
try:
    from local_settings import *       # noqa: F401,F403
except ImportError:
    pass

import argparse
import base64
import json
import os
import re
import shutil
import sys
import webbrowser
from datetime import datetime
from pathlib import Path

try:
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Python packages missing. Install them once with:\n    py -m pip install pandas openpyxl\nthen run again.")

HERE = Path(__file__).resolve().parent
ORANGE, CHARCOAL, PEACH = "D04A02", "2D2D2D", "FDEDE4"

# ------------------------------------------------------------ workbook layout ---
SHEET_RAW = "Project Hours Raw Data"
SHEET_STAFF = "Staff Master"
SHEET_MODEL = "Global Model FY27 CPP"
SHEET_CALLOUTS = "Callouts"
SHEET_ASSUMPTIONS = "Assumptions"
SHEET_WRN_CAT = "WRN Category"
# Optional input tabs (static content you maintain). Missing tabs fall back to the defaults in the scripts.
SHEET_TAG_MAP = "Tag Mapping"
SHEET_TEAM = "Team"
SHEET_SERVICES = "Services"
SHEET_THEMES = "ET Themes"
SHEET_STORY = "Story"
SHEET_PRESETS = "Scenario Presets"
# Used only if the workbook has no 'Global Model FY27 CPP' sheet.
# Illustrative values only - put your real rate card in the workbook.
DEFAULT_MODEL = {"SM": ("Senior Manager", 95, 96000), "M": ("Manager", 65, 64000),
                 "SA": ("Senior Associate", 48, 43000), "A": ("Associate", 30, 28500)}
#   # optional: Work Request No. -> ESG/ET, fills blanks in the raw ESG/ET column
CAT_COLUMN = "ESG/ET"

# Display order; any other territory found in the data is appended after these.
TERRITORY_ORDER = ["Canada", "USA", "Australia", "UK", "Japan"]
# Spelling variants in the raw data -> standard name.
TERRITORY_ALIASES = {
    "us": "USA", "usa": "USA", "united states": "USA",
    "uk": "UK", "united kingdom": "UK",
    "au": "Australia", "aus": "Australia", "australia": "Australia",
    "ca": "Canada", "canada": "Canada",
    "jp": "Japan", "japan": "Japan",
}
# Keyword in 'Home Sub LOS 2' -> territory the person is currently tagged to.
# Tags such as 'Other Territory' or 'Deals' map to no territory.
TAG_KEYWORDS = {r"\bUS\b|\bUSA\b": "USA", r"\bUK\b": "UK", r"\bCA\b|Canada": "Canada",
                r"\bAU\b|Australia": "Australia", r"\bJP\b|Japan": "Japan"}

# Scenario presets shown as one-click buttons on the 'US tagging scenario' tab.
# threshold / window = None -> use the value from the Assumptions sheet.
# growth = expected change in US hours next year (0.10 = +10%).
# method: "hourly" = Global Model rate/hr; "best" = cheaper of hourly or a dedicated seat (Cost/Person).
SCENARIO_PRESETS = [
    {"name": "Conservative", "threshold": 0.75, "window": 12, "growth": 0.00, "method": "hourly"},
    {"name": "Base", "threshold": None, "window": None, "growth": 0.10, "method": "hourly"},
    {"name": "Growth", "threshold": 0.50, "window": 12, "growth": 0.30, "method": "best"},
]

DEFAULT_TAGGING_NOTE = ("Today no one is tagged to a territory: everyone sits under a non-territory tag such as "
                        "Other Territory. This analysis prepares the move to territory tagging.")

DEFAULT_ASSUMPTIONS = {"FY start month": 7, "Tagging threshold": 0.6,
                       "Scenario window (months)": 12, "Annual capacity hours": 1800,
                       "Training month": "Aug-24"}

# ------------------------------------------------- default text for the input tabs ---
# Used only to pre-fill tabs that are missing from a workbook. After that, edit the text in Excel.
STORY = {
    "title": "Tagging US-focused AC talent to the US",
    "subtitle": "A win-win proposal: the US pays less for the same people, and AC builds a dedicated US bench",
    "meeting": "Discussion with US Partners",
    "presenter": "Your Team, Your Organisation",
    "ask": [
        "Approve tagging to the US of the AC people who already spend most of their time on US work",
        "Agree the Global Model FY27 rate card as the basis for US-tagged hours",
        "Review the tagging every quarter as US volume grows, and add people as they cross the threshold",
    ],
    "winUS": [
        "Lower cost for the same people and the same quality: Global Model rate instead of sales price",
        "The yearly saving can fund extra capacity, such as an additional Associate",
        "A named, dedicated US team that knows US clients, frameworks and ways of working",
        "Hours and cost show up correctly in US reporting",
    ],
    "winAC": [
        "Stable, predictable US demand to plan hiring and skills around",
        "People are tagged to where their work actually is: cleaner utilisation and P&L",
        "Deeper US specialisation and career paths for the team",
        "A model we can repeat for UK, Canada, Australia and Japan",
    ],
    "steps": [
        ["Now", "Tag the people above the threshold to the US", "Confirm the rate card with Finance; update Home Sub LOS 2 for the agreed people"],
        ["Next quarter", "Review shared and growing profiles", "People whose US share passes the threshold move next; track the saving monthly"],
        ["FY28", "Dedicated US pod", "Move full-time US people onto dedicated seats (Cost/Person) where that is cheaper"],
    ],
}

DEFAULT_CONTENT = {'team': {'teamName': 'Your Team Name',
          'summary': 'Team supports global Deals through research, analysis, and due diligence across '
                     'Sustainability/ESG/Energy Transition/EUR/IUR topics',
          'leaders': [{'name': 'Leader One', 'role': 'Managing Director', 'note': 'Practice Leader'},
                      {'name': 'Leader Two', 'role': 'Director', 'note': ''},
                      {'name': 'Leader Three', 'role': 'Manager', 'note': 'Primary point of contact'}],
          'members': '6 Members',
          'keyContacts': [{'territory': 'Australia', 'name': 'Contact Name', 'role': 'Partner'},
                          {'territory': 'USA', 'name': 'Contact Name', 'role': 'Partner'},
                          {'territory': 'Japan', 'name': 'Contact Name', 'role': 'Partner'},
                          {'territory': 'Canada', 'name': 'Contact Name', 'role': 'Director'}],
          'currentServices': ['Company Profiles/Target Snapshot',
                              'Industry ESG trends and analysis',
                              'Industry ESG regulatory overview',
                              'Value chain analysis',
                              'Customer commitments and targets analysis',
                              'Comparable company list',
                              'ESG materiality assessment',
                              'ESG maturity assessment',
                              'ESG metrics benchmarking',
                              'Social media and reputational risk scan',
                              'Industry ESG trends repository',
                              'Industry ESG regulations repository',
                              'Drafting Management interview questions',
                              'Sentiment analysis',
                              'Virtual Data Room data review'],
          'futureServices': ['Drafting Request for information (RFI)',
                             'ESG metrics benchmark repository',
                             'ESG Framework & Disclosures repository',
                             'Red Flags Screening',
                             'ESG Due Diligence (first draft)',
                             'ESG Action Plan (first draft)',
                             'ESG Value Creation Plan',
                             'High level Climate Risk assessment',
                             'Thought Leadership Support',
                             'Proposal Support'],
          'frameworks': ['SASB Standards', 'MSCI - ESG Industry Materiality Map', 'PRI', 'CDP'],
          'databases': ['Pitchbook',
                        'Merger Market',
                        'S&P Capital IQ',
                        'Inframation',
                        'Statista',
                        'Technavio',
                        'Factiva',
                        'Bloomberg',
                        'IBIS World (US)',
                        'Debtwire',
                        'Lexis Nexis (only for US)'],
          'aiTools': ['Web Plot Digitizer',
                      'Copilot',
                      'Cursor',
                      'Claude Design',
                      'Internal GenAI assistant',
                      'ChatGPT Enterprise',
                      'Replit',
                      'Harvey',
                      'Tolt',
                      'Visual Studio Code']},
 'themes': {'title': 'Our Work Across Key Energy Transition Themes',
            'items': [{'title': 'Clean Energy Systems',
                       'tagline': 'Core technologies enabling low-carbon transitions',
                       'points': [['Hydrogen Mobility & Fuel',
                                   'Market trends analysis across Asia; vehicle adoption studies by global OEMs'],
                                  ['EV Ecosystem',
                                   'Market assessment for EV deals; EV charging infrastructure trends and policy '
                                   'mapping'],
                                  ['Battery Energy Storage (BESS)',
                                   'Research on Australian BESS projects (stage, size, distribution)'],
                                  ['Renewable Energy Assets',
                                   'Analysis of 600+ renewable projects (capex, equity ownership, conflict risks)'],
                                  ['Sustainable Aviation Fuel (SAF)',
                                   'R&D on biofuel-based jet fuel pathways vs conventional fuel']]},
                      {'title': 'Carbon Markets, Climate Governance & ESG Frameworks',
                       'tagline': 'Enabling credibility, compliance & climate governance maturity',
                       'points': [['TCFD Framework', 'Slide preparation, peer benchmarking & gap analysis'],
                                  ['Voluntary Carbon Markets (VCM)',
                                   'Countrywide policy analysis; comparison of IC-VCM/VCMI standards'],
                                  ['Biodiversity Credit Systems',
                                   'Research on nature-based credit mechanisms supporting net zero'],
                                  ['ESG Frameworks & Standards', 'Benchmarking ESG disclosures & commitments'],
                                  ['Policy & Regulatory Analysis',
                                   'Governance structures for renewable energy and carbon markets'],
                                  ['Carbon Capture & Storage (CCS)', 'Multi-country project assessments']]},
                      {'title': 'Decarbonisation Pathways & Sector Transitions',
                       'tagline': 'Supporting industries in developing credible net-zero strategies',
                       'points': [['Mine Decarbonisation', 'Benchmarking commitments; energy source transition plans'],
                                  ['Real Estate Transition', 'Mapping ESG KPIs & transition strategies'],
                                  ['Urban Regeneration Projects',
                                   'Analysis of low-carbon infrastructure options (EV charging, green mobility)'],
                                  ['Education Sector Decarbonisation',
                                   'Sustainability journey assessment & net-zero roadmap'],
                                  ['Industrial Emission Reduction',
                                   'Options assessment across infrastructure & built environment'],
                                  ['DERMs & VPP Engineering', 'Global competitor and business model analysis']]},
                      {'title': 'Circular Economy, Resource Efficiency & Environmental Solutions',
                       'tagline': 'Reducing waste, improving resource use & supporting nature-positive outcomes',
                       'points': [['Solar Panel Recycling', 'Market study of obsolete panel recycling in Australia'],
                                  ['Waste Management',
                                   'Company profiling, project analysis & government grant mapping'],
                                  ['Water & Wastewater Treatment',
                                   'Project identification & future potential assessment'],
                                  ['Critical Minerals', 'Screening key projects and gathering intelligence'],
                                  ['Environment, Conservation & Wildlife',
                                   'Strategic ESG guidance for environmental outcomes'],
                                  ['Investor Participation in Renewables',
                                   'Investor landscape analysis & investment outlook']]}]}}


# ================================================================ read the data ===
def norm_territory(v):
    s = str(v).strip()
    return TERRITORY_ALIASES.get(s.lower(), s)


def to_month(v):
    """'Sep-24', a date, or 'Sep 2024' -> Period('2024-09')."""
    if pd.isna(v) or str(v).strip() == "":
        return None
    if isinstance(v, (datetime, pd.Timestamp)):
        return pd.Period(v, "M")
    s = str(v).strip()
    for fmt in ("%b-%y", "%b-%Y", "%b %y", "%b %Y", "%Y-%m", "%m/%Y"):
        try:
            return pd.Period(datetime.strptime(s, fmt), "M")
        except ValueError:
            pass
    try:
        return pd.Period(pd.to_datetime(s, dayfirst=True), "M")
    except Exception:
        return None


def clean_str(v, default=""):
    return default if pd.isna(v) or str(v).strip() == "" else str(v).strip()


def load(excel_path):
    xl = pd.ExcelFile(excel_path)
    need = [SHEET_RAW]
    missing = [s for s in need if s not in xl.sheet_names]
    if missing:
        sys.exit(f"Missing sheet(s): {missing}. Found: {xl.sheet_names}")
    raw = xl.parse(SHEET_RAW)
    staff = xl.parse(SHEET_STAFF) if SHEET_STAFF in xl.sheet_names else pd.DataFrame()
    model = xl.parse(SHEET_MODEL) if SHEET_MODEL in xl.sheet_names else pd.DataFrame()
    callouts = xl.parse(SHEET_CALLOUTS) if SHEET_CALLOUTS in xl.sheet_names else pd.DataFrame()
    assum = dict(DEFAULT_ASSUMPTIONS)
    if SHEET_ASSUMPTIONS in xl.sheet_names:
        a = xl.parse(SHEET_ASSUMPTIONS)
        for _, r in a.iterrows():
            if clean_str(r.get("Parameter")):
                assum[clean_str(r["Parameter"])] = r["Value"]
    wrn_cat = xl.parse(SHEET_WRN_CAT, dtype=str) if SHEET_WRN_CAT in xl.sheet_names else pd.DataFrame()
    inputs = {name: xl.parse(name).rename(columns=lambda c: str(c).strip()).dropna(how="all")
              for name in (SHEET_TAG_MAP, SHEET_TEAM, SHEET_SERVICES, SHEET_THEMES, SHEET_STORY, SHEET_PRESETS)
              if name in xl.sheet_names}
    return raw, staff, model, callouts, assum, wrn_cat, inputs


def to_share(v, default=None):
    """60, '60%', 0.6 -> 0.6 ; blank -> default"""
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "":
        return default
    s = str(v).strip().replace("%", "")
    try:
        x = float(s)
    except ValueError:
        return default
    return x / 100 if abs(x) > 1.5 or str(v).strip().endswith("%") else x


def read_inputs(inputs, assum):
    """Turn the optional input tabs into: tag map, page content, story text, scenario presets."""
    out = {"tagMap": {}, "content": None, "story": None, "presets": None}
    df = inputs.get(SHEET_TAG_MAP)
    if df is not None and "Home Sub LOS 2" in df.columns:
        for tag, terr in zip(df["Home Sub LOS 2"], df.get("Territory", [""] * len(df))):
            if clean_str(tag):
                out["tagMap"][clean_str(tag).lower()] = norm_territory(terr) if clean_str(terr) else ""

    team, svc, th = inputs.get(SHEET_TEAM), inputs.get(SHEET_SERVICES), inputs.get(SHEET_THEMES)
    if team is not None or svc is not None or th is not None:
        c = {"team": {"teamName": "", "summary": "", "members": "", "leaders": [], "keyContacts": [], "currentServices": [],
                      "futureServices": [], "frameworks": [], "databases": [], "aiTools": []},
             "themes": {"title": clean_str(assum.get("Energy Transition page title"), "Our Work Across Key Energy Transition Themes"), "items": []}}
        t = c["team"]
        if team is not None:
            for _, r in team.iterrows():
                typ = clean_str(r.get("Type")).lower()
                nm, rl, tr, nt = (clean_str(r.get(k)) for k in ("Name", "Role", "Territory", "Note"))
                if typ == "team name": t["teamName"] = nm
                elif typ == "summary": t["summary"] = nt or nm
                elif typ == "members": t["members"] = nm or nt
                elif typ == "leader": t["leaders"].append({"name": nm, "role": rl, "note": nt})
                elif typ == "key contact": t["keyContacts"].append({"territory": norm_territory(tr), "name": nm, "role": rl})
        if svc is not None:
            keymap = {"current service": "currentServices", "future service": "futureServices", "esg framework": "frameworks",
                      "database": "databases", "digital & ai tool": "aiTools", "ai tool": "aiTools"}
            for cat, item in zip(svc.get("Category", []), svc.get("Item", [])):
                k = keymap.get(clean_str(cat).lower())
                if k and clean_str(item): t[k].append(clean_str(item))
        if th is not None:
            for no, g in th.groupby(th["Theme No"].ffill(), sort=True):
                first = g.iloc[0]
                c["themes"]["items"].append({"title": clean_str(first.get("Theme Title")), "tagline": clean_str(first.get("Tagline")),
                    "points": [[clean_str(a), clean_str(b)] for a, b in zip(g.get("Work Area", []), g.get("Description", [])) if clean_str(a)]})
        out["content"] = c

    st = inputs.get(SHEET_STORY)
    if st is not None and "Section" in st.columns:
        story = {}
        lists = {"ask": "ask", "win for us": "winUS", "win for ac": "winAC"}
        for _, r in st.iterrows():
            sec = clean_str(r.get("Section")).lower(); tx = clean_str(r.get("Text"))
            if sec in ("title", "subtitle", "meeting", "presenter"): story[sec] = tx
            elif sec in lists: story.setdefault(lists[sec], []).append(tx)
            elif sec == "next step": story.setdefault("steps", []).append([tx, clean_str(r.get("Heading")), clean_str(r.get("Detail"))])
        out["story"] = story

    pr = inputs.get(SHEET_PRESETS)
    if pr is not None and "Name" in pr.columns:
        out["presets"] = [{"name": clean_str(r.get("Name")),
                           "threshold": to_share(r.get("Tagging threshold")),
                           "window": int(r["Look-back months"]) if pd.notna(r.get("Look-back months")) else None,
                           "growth": to_share(r.get("US volume change next year"), 0.0),
                           "method": "best" if "seat" in clean_str(r.get("Costing method")).lower() else "hourly"}
                          for _, r in pr.iterrows() if clean_str(r.get("Name"))]
    return out


def build_payload(raw, staff, model, callouts, assum, wrn_cat, inputs, source_name):
    inp = read_inputs(inputs or {}, assum)
    raw = raw.rename(columns=lambda c: str(c).strip())
    for col in ["Territory Name", "Staff Name", "Hours"]:
        if col not in raw.columns:
            sys.exit(f"Column '{col}' not found in '{SHEET_RAW}'.")

    # Month: 'Source Sheet' is the reporting month ("consider as date"); fall back to Date.
    month = raw["Source Sheet"].map(to_month) if "Source Sheet" in raw.columns else pd.Series([None] * len(raw))
    if "Date" in raw.columns:
        month = month.where(month.notna(), raw["Date"].map(to_month))
    raw["_month"] = month
    raw["Hours"] = pd.to_numeric(raw["Hours"], errors="coerce")
    bad = raw["_month"].isna() | raw["Hours"].isna() | raw["Staff Name"].isna() | raw["Territory Name"].isna()
    dropped = int(bad.sum())
    raw = raw[~bad].copy()
    raw["_terr"] = raw["Territory Name"].map(norm_territory)
    raw["_staff"] = raw["Staff Name"].map(clean_str)

    # Staff master
    smap = {}
    rate_cols = {}
    if not staff.empty:
        staff = staff.rename(columns=lambda c: str(c).strip())
        for c in staff.columns:
            m = re.search(r"(?:Sales Price|Rate)\D*(\d{4})", c, re.I)
            if m:
                rate_cols[int(m.group(1))] = c
        plain = [c for c in staff.columns if re.fullmatch(r"(sales price|rate)\b.*", c, re.I) and not re.search(r"\d{4}", c)]
        if not rate_cols and plain:   # one 'Sales Price' column: use it for every year in the data
            rate_cols = {y: plain[0] for y in sorted({p.year for p in raw["_month"]})}
        for _, r in staff.iterrows():
            name = clean_str(r.get("Employee"))
            if not name:
                continue
            tag_txt = clean_str(r.get("Home Sub LOS 2"))
            tag_terr = inp["tagMap"].get(tag_txt.lower())
            for pat, terr in ([] if tag_terr is not None else TAG_KEYWORDS.items()):
                if re.search(pat, tag_txt, re.I):
                    tag_terr = terr
                    break
            tag_terr = tag_terr or ""
            smap[name] = {
                "name": name,
                "level": clean_str(r.get("Level")).upper(),
                "team": clean_str(r.get("Team Name")),
                "costCenter": clean_str(r.get("Home Cost Center")),
                "tag": tag_txt,
                "tagTerritory": tag_terr,
                "rates": {str(y): float(r[c]) for y, c in rate_cols.items() if pd.notna(r.get(c))},
            }

    # Fill level / team from the raw data where the master lacks them, and vice versa.
    unknown_staff = sorted(set(raw["_staff"]) - set(smap))
    for name in unknown_staff:
        smap[name] = {"name": name, "level": "", "team": "", "costCenter": "", "tag": "", "tagTerritory": "", "rates": {}}
    if "Management Level" in raw.columns:
        for name, lv in raw.dropna(subset=["Management Level"]).groupby("_staff")["Management Level"].agg(lambda s: s.mode().iat[0]).items():
            if not smap[name]["level"]:
                smap[name]["level"] = str(lv).strip().upper()
    if "Team Name" in raw.columns:
        raw["_team"] = raw["Team Name"].map(clean_str)
        for name, tm in raw[raw["_team"] != ""].groupby("_staff")["_team"].agg(lambda s: s.mode().iat[0]).items():
            if not smap[name]["team"]:
                smap[name]["team"] = tm
        raw.loc[raw["_team"] == "", "_team"] = raw["_staff"].map(lambda n: smap[n]["team"])
    else:
        raw["_team"] = raw["_staff"].map(lambda n: smap[n]["team"])
    raw["_team"] = raw["_team"].replace("", "Unassigned")

    # Projects
    pcols = {"name": "Project Name", "client": "Client Name/WBS Name", "wbs": "Client code/WBS code",
             "wrn": "Work Request No.", "job": "Job Name"}
    for k, c in pcols.items():
        raw["_p_" + k] = raw[c].map(clean_str) if c in raw.columns else ""
    raw.loc[raw["_p_name"] == "", "_p_name"] = raw["_p_job"].where(raw["_p_job"] != "", "Unspecified")
    raw["_pkey"] = raw["_terr"] + "|" + raw["_p_name"] + "|" + raw["_p_client"]
    raw["_flag"] = raw["Project Match Flag"].map(lambda v: clean_str(v, "Not flagged")) if "Project Match Flag" in raw.columns else "Not flagged"

    # ESG / ET category: the row's own ESG/ET value, else the WRN Category sheet, else 'Unclassified'
    norm_cat = lambda v: {"ENERGY TRANSITION": "ET"}.get(clean_str(v).upper(), clean_str(v).upper())
    raw["_cat"] = raw[CAT_COLUMN].map(norm_cat) if CAT_COLUMN in raw.columns else ""
    if not wrn_cat.empty and {"Work Request No.", CAT_COLUMN} <= set(wrn_cat.columns):
        lookup = {clean_str(w): norm_cat(c) for w, c in zip(wrn_cat["Work Request No."], wrn_cat[CAT_COLUMN]) if clean_str(c)}
        blank = raw["_cat"] == ""
        raw.loc[blank, "_cat"] = raw.loc[blank, "_p_wrn"].map(lambda w: lookup.get(w, ""))
    raw.loc[raw["_cat"] == "", "_cat"] = "Unclassified"

    # Index tables
    months = pd.period_range(raw["_month"].min(), raw["_month"].max(), freq="M")
    m_idx = {p: i for i, p in enumerate(months)}
    terrs = [t for t in TERRITORY_ORDER if t in set(raw["_terr"])] + sorted(set(raw["_terr"]) - set(TERRITORY_ORDER))
    t_idx = {t: i for i, t in enumerate(terrs)}
    staff_names = sorted(smap)
    s_idx = {n: i for i, n in enumerate(staff_names)}
    teams = sorted(set(raw["_team"]))
    tm_idx = {t: i for i, t in enumerate(teams)}
    flags = sorted(set(raw["_flag"]))
    f_idx = {f: i for i, f in enumerate(flags)}
    pdf = raw.drop_duplicates("_pkey")
    projects = [{"territory": r["_terr"], "name": r["_p_name"], "client": r["_p_client"], "wbs": r["_p_wbs"],
                 "wrn": r["_p_wrn"], "job": r["_p_job"]} for _, r in pdf.iterrows()]
    p_idx = {k: i for i, k in enumerate(pdf["_pkey"])}
    pcat = raw.groupby("_pkey")["_cat"].agg(lambda x: x.mode().iat[0])
    for k, i in p_idx.items():
        projects[i]["cat"] = pcat[k]
    cats = [c for c in ["ESG", "ET"] if c in set(raw["_cat"])] + sorted(set(raw["_cat"]) - {"ESG", "ET"})
    c_idx = {c: i for i, c in enumerate(cats)}

    rows = [[m_idx[m], t_idx[t], s_idx[s], p_idx[p], tm_idx[tm], round(float(h), 2), f_idx[f], c_idx[c]]
            for m, t, s, p, tm, h, f, c in zip(raw["_month"], raw["_terr"], raw["_staff"], raw["_pkey"],
                                               raw["_team"], raw["Hours"], raw["_flag"], raw["_cat"])]

    model_rates = {}
    if not model.empty:
        model = model.rename(columns=lambda c: str(c).strip())
        for _, r in model.iterrows():
            lv = clean_str(r.get("Level")).upper()
            if lv:
                model_rates[lv] = {"name": clean_str(r.get("Level Name"), lv),
                                   "rate": float(r.get("Rate/hr", 0) or 0),
                                   "cpp": float(r.get("Cost/Person", 0) or 0)}

    model_default = not model_rates
    if model_default:
        model_rates = {lv: {"name": n, "rate": float(r), "cpp": float(c)} for lv, (n, r, c) in DEFAULT_MODEL.items()}

    calls = []
    if not callouts.empty:
        for _, r in callouts.iterrows():
            p = to_month(r.get("Month"))
            if p is not None:
                calls.append({"month": str(p), "territory": clean_str(r.get("Territory")), "event": clean_str(r.get("Event"))})

    tm = to_month(assum.get("Training month"))
    meta = {
        "source": source_name,
        "generated": datetime.now().strftime("%d %b %Y %H:%M"),
        "fyStart": int(assum.get("FY start month", 7)),
        "threshold": float(assum.get("Tagging threshold", 0.6)),
        "window": int(assum.get("Scenario window (months)", 12)),
        "capacity": float(assum.get("Annual capacity hours", 1800)),
        "trainingMonth": str(tm) if tm is not None else None,
        "taggingNote": clean_str(assum.get("Current tagging"), DEFAULT_TAGGING_NOTE),
        "rateYears": [str(y) for y in sorted(set(rate_cols) & {p.year for p in months})][-2:],
        "title": clean_str(assum.get("Report title"), REPORT_TITLE), "classification": clean_str(assum.get("Classification"), CLASSIFICATION),
        "contact": {"people": [{"name": n, "email": e} for n, e in ((CONTACT_NAME, CONTACT_EMAIL), (CONTACT_NAME_2, CONTACT_EMAIL_2)) if n or e],
                    "team": CONTACT_TEAM},
        "storyLink": STORY_FILE,
    }
    meta["presets"] = [{**pr, "threshold": meta["threshold"] if pr["threshold"] is None else pr["threshold"],
                        "window": meta["window"] if pr["window"] is None else pr["window"]} for pr in (inp["presets"] or SCENARIO_PRESETS)]
    if not any(p["name"].lower() == "base" for p in meta["presets"]) and meta["presets"]:
        meta["presets"][0]["name"] = meta["presets"][0]["name"] or "Base"
    payload = {"meta": meta, "months": [str(p) for p in months], "territories": terrs,
               "staff": [smap[n] for n in staff_names], "teams": teams, "flags": flags, "categories": cats,
               "projects": projects, "rows": rows, "model": model_rates, "callouts": calls,
               "content": inp["content"], "storyInput": inp["story"]}

    report = {
        "rows used": len(rows), "rows dropped (no month/hours/staff/territory)": dropped,
        "total hours": round(raw["Hours"].sum(), 1),
        "hours by ESG/ET": ", ".join(f"{c} {h:,.0f}" for c, h in raw.groupby("_cat")["Hours"].sum().items()),
        "months": f"{months[0]} to {months[-1]}", "territories": ", ".join(terrs),
        "staff not in Staff Master": ", ".join(unknown_staff) or "none",
        "staff without a level": ", ".join(n for n in staff_names if not smap[n]["level"]) or "none",
        "sales prices": (", ".join(f"{y}: '{c}'" for y, c in sorted(rate_cols.items())) if rate_cols
                         else "MISSING - add 'Sales Price' (or 'Sales Price YYYY') to Staff Master; cost and saving figures need it"),
        "rate card": "Global Model sheet" if not model_default else "sheet missing - using built-in FY27 defaults",
        "input tabs found": ", ".join(inputs) if inputs else "none (using defaults in the scripts)",
        "levels missing from Global Model": ", ".join(sorted({smap[n]["level"] for n in staff_names if smap[n]["level"] and smap[n]["level"] not in model_rates})) or "none",
    }
    return payload, report



# ============================================================ input tabs in Excel ===
def tables(wb):
    """Return {sheet name: (rows, column widths, note)} for every input tab."""
    team, themes = DEFAULT_CONTENT["team"], DEFAULT_CONTENT["themes"]
    story = STORY
    out = {}

    # Tag Mapping: every Home Sub LOS 2 value found in Staff Master
    tags = []
    if SHEET_STAFF in wb.sheetnames:
        sm = pd.DataFrame(wb[SHEET_STAFF].values)
        if len(sm):
            sm.columns = [str(c).strip() for c in sm.iloc[0]]; sm = sm.iloc[1:]
            if "Home Sub LOS 2" in sm.columns:
                tags = sorted({str(v).strip() for v in sm["Home Sub LOS 2"].dropna() if str(v).strip()})
    def guess(tag):
        for pat, terr in TAG_KEYWORDS.items():
            if re.search(pat, tag, re.I):
                return terr
        return ""
    out[SHEET_TAG_MAP] = ([["Home Sub LOS 2", "Territory"]] + [[t, guess(t)] for t in tags],
                             [36, 16], "Which territory each Home Sub LOS 2 tag means. Leave Territory blank for tags that are not a territory (e.g. Other Territory, Deals).")

    rows = [["Type", "Name", "Role", "Territory", "Note"],
            ["Team name", team.get("teamName", ""), "", "", ""],
            ["Summary", "", "", "", team.get("summary", "")],
            ["Members", team.get("members", ""), "", "", ""]]
    rows += [["Leader", l.get("name", ""), l.get("role", ""), "", l.get("note", "")] for l in team.get("leaders", [])]
    rows += [["Key contact", k.get("name", ""), k.get("role", ""), k.get("territory", ""), ""] for k in team.get("keyContacts", [])]
    out[SHEET_TEAM] = (rows, [14, 26, 20, 14, 70], "Type must be one of: Team name, Summary, Members, Leader, Key contact. Add or delete Leader / Key contact rows freely.")

    rows = [["Category", "Item"]]
    for cat, key in [("Current service", "currentServices"), ("Future service", "futureServices"), ("ESG framework", "frameworks"),
                     ("Database", "databases"), ("Digital & AI tool", "aiTools")]:
        rows += [[cat, x] for x in team.get(key, [])]
    out[SHEET_SERVICES] = (rows, [20, 50], "Category must be one of: Current service, Future service, ESG framework, Database, Digital & AI tool.")

    rows = [["Theme No", "Theme Title", "Tagline", "Work Area", "Description"]]
    for i, it in enumerate(themes.get("items", []), 1):
        for j, (a, d) in enumerate(it.get("points", [])):
            rows.append([i, it["title"] if j == 0 else "", it["tagline"] if j == 0 else "", a, d])
    out[SHEET_THEMES] = (rows, [10, 34, 40, 30, 60], "One row per work area. Theme Title and Tagline only need filling on the first row of each theme. Up to 4 themes are shown.")

    rows = [["Section", "Text", "Heading", "Detail"]]
    for k in ("title", "subtitle", "meeting", "presenter"):
        if story.get(k): rows.append([k.capitalize(), story[k], "", ""])
    rows += [["Ask", x, "", ""] for x in story.get("ask", [])]
    rows += [["Win for US", x, "", ""] for x in story.get("winUS", [])]
    rows += [["Win for AC", x, "", ""] for x in story.get("winAC", [])]
    rows += [["Next step", w, t, d] for w, t, d in story.get("steps", [])]
    out[SHEET_STORY] = (rows, [14, 70, 34, 60], "Section: Title, Subtitle, Meeting, Presenter, Ask, Win for US, Win for AC, Next step. For Next step, Text = when (e.g. 'Now').")

    rows = [["Name", "Tagging threshold", "Look-back months", "US volume change next year", "Costing method"]]
    for p in SCENARIO_PRESETS:
        rows.append([p["name"], p["threshold"], p["window"], p["growth"], "Hourly or seat" if p["method"] == "best" else "Hourly"])
    out[SHEET_PRESETS] = (rows, [16, 18, 18, 26, 18], "Blank threshold / look-back = use the Assumptions tab. Costing method: Hourly, or Hourly or seat.")
    return out


README = [
    ["Tab", "What it is", "Who updates it", "How"],
    [SHEET_RAW, "Time entries (one row per entry)", "Monthly export", "Paste the new export; keep the column headers"],
    [SHEET_STAFF, "People: level, team, Home Sub LOS 2 tag, sales price", "When people join, move or rates change", "One row per person; 'Sales Price' or 'Sales Price 2025', 'Sales Price 2026'..."],
    [SHEET_MODEL, "Global Model rate card (Rate/hr, Cost/Person by level)", "Once a year", "Levels must match Staff Master (SM, M, SA, A...)"],
    [SHEET_CALLOUTS, "Milestones shown on the monthly chart", "As needed", "Month as Mon-YY, e.g. Feb-25"],
    [SHEET_ASSUMPTIONS, "Settings: FY start, tagging threshold, look-back, titles", "Rarely", "Change the Value column only"],
    [SHEET_WRN_CAT, "ESG or ET per Work Request No.", "When new WRNs appear", "Fills blank ESG/ET cells in the raw data"],
    [SHEET_TAG_MAP, "Which territory each Home Sub LOS 2 tag means", "When a new tag appears", "Blank Territory = not a territory tag"],
    [SHEET_TEAM, "Team overview: leaders, members, key contacts", "As the team changes", "See the note on the tab"],
    [SHEET_SERVICES, "Services, frameworks, databases, AI tools", "As needed", "See the note on the tab"],
    [SHEET_THEMES, "Energy Transition themes page", "As needed", "See the note on the tab"],
    [SHEET_STORY, "Text of the partner story page", "Before each meeting", "See the note on the tab"],
    [SHEET_PRESETS, "Scenario presets (Conservative / Base / Growth)", "As needed", "See the note on the tab"],
    ["", "", "", ""],
    ["After editing", "Save and close the workbook, then run:  py build.py   (or double-click build.py)", "", ""],
]
EXTRA_ASSUMPTIONS = [
    ("Report title", REPORT_TITLE, "Title at the top of the dashboard."),
    ("Classification", CLASSIFICATION, "Shown in the header and footer of both pages."),
    ("Energy Transition page title", "Our Work Across Key Energy Transition Themes", "Heading of the Energy Transition tab."),
    ("Current tagging", DEFAULT_TAGGING_NOTE, "Shown in the assumptions on both pages. Update once people are tagged to territories."),
]


def style_sheet(ws, widths, header_fill=CHARCOAL):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=header_fill)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"



def ensure_input_tabs(path):
    """Add any missing input tabs (pre-filled). Saves only if something was added; keeps one backup."""
    wb = load_workbook(path)
    added = []

    for name, (rows, widths, note) in tables(wb).items():
        if name in wb.sheetnames:
            continue
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        style_sheet(ws, widths, ORANGE)
        ws.sheet_properties.tabColor = ORANGE
        nc = ws.cell(row=1, column=len(rows[0]) + 2, value="How to fill: " + note)   # beside the table, so it is never read as data
        nc.font = Font(italic=True, color="7D7D7D"); nc.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(len(rows[0]) + 2)].width = 48
        added.append(name)

    # Assumptions: add the title rows if missing
    if SHEET_ASSUMPTIONS in wb.sheetnames:
        ws = wb[SHEET_ASSUMPTIONS]
        have = {str(r[0].value).strip() for r in ws.iter_rows(min_row=2) if r[0].value}
        for p, v, n in EXTRA_ASSUMPTIONS:
            if p not in have:
                ws.append([p, v, n]); added.append(f"Assumptions: {p}")
    else:
        ws = wb.create_sheet(SHEET_ASSUMPTIONS)
        ws.append(["Parameter", "Value", "Notes"])
        for p, v in DEFAULT_ASSUMPTIONS.items():
            ws.append([p, v, ""])
        for p, v, n in EXTRA_ASSUMPTIONS:
            ws.append([p, v, n])
        style_sheet(ws, [30, 40, 70]); added.append(SHEET_ASSUMPTIONS)

    if "README" not in wb.sheetnames:
        ws = wb.create_sheet("README", 0)
        for r in README:
            ws.append(r)
        style_sheet(ws, [26, 52, 30, 70], ORANGE)
        ws.sheet_properties.tabColor = CHARCOAL
        added.append("README")

    if not added:
        print("Input tabs: all present, workbook not changed")
        return
    backup = path.with_name(path.stem + " (backup before input tabs)" + path.suffix)
    if not backup.exists():
        shutil.copy2(path, backup)     # keep the original once, in case Excel features need restoring
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Could not save {path.name} - close it in Excel and run again.")
    print("Input tabs added:", ", ".join(added), f"(original kept as '{backup.name}')")



# ================================================================== write pages ===
def load_templates():
    """template.html holds both pages and the chart library, separated by <!-- @@NAME --> markers."""
    path = HERE / TEMPLATE_FILE
    if not path.exists():
        sys.exit(f"{TEMPLATE_FILE} not found next to build.py.")
    parts = re.split(r"<!-- @@([A-Z]+)[^>]*-->\n?", path.read_text(encoding="utf-8"))
    found = {name: body for name, body in zip(parts[1::2], parts[2::2])}
    for need in ("DASHBOARD", "STORY", "CHARTJS"):
        if need not in found:
            sys.exit(f"{TEMPLATE_FILE} is missing the <!-- @@{need} --> section.")
    return found


def logo_html():
    logo = Path(LOGO_FILE) if Path(LOGO_FILE).is_absolute() else HERE / LOGO_FILE
    if not logo.exists():
        return f'<span class="wordmark">{LOGO_TEXT}</span>'
    mime = {".svg": "image/svg+xml", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(logo.suffix.lower(), "image/png")
    return f'<img class="logo" alt="logo" src="data:{mime};base64,{base64.b64encode(logo.read_bytes()).decode()}">'


def write_page(template, chart_tag, payload, out):
    html = (template.replace("/*__CHARTJS__*/", chart_tag).replace("<!--__LOGO__-->", logo_html())
                    .replace("/*__DATA__*/null", json.dumps(payload, separators=(",", ":"))))
    out.write_text(html, encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--excel", default=EXCEL_FILE, help="Path to the hours workbook")
    ap.add_argument("--out-dir", default=str(HERE), help="Folder for the two HTML pages")
    ap.add_argument("--no-open", action="store_true", help="Do not open the story page at the end")
    args = ap.parse_args()
    excel = Path(args.excel).expanduser()
    excel = excel if excel.is_absolute() else HERE / excel
    if not excel.exists():
        sys.exit(f"Excel file not found: {excel}\nSet EXCEL_FILE at the top of build.py (or in local_settings.py).")
    out_dir = Path(args.out_dir).expanduser(); out_dir.mkdir(parents=True, exist_ok=True)
    tpl = load_templates()

    print(f"Workbook: {excel.name}\n\n[1/3] Checking input tabs ...")
    ensure_input_tabs(excel)

    print("\n[2/3] Reading the data ...")
    payload, report = build_payload(*load(excel), source_name=excel.name)
    if payload["content"] is None:
        payload["content"] = DEFAULT_CONTENT
    payload["story"] = {**STORY, **{k: v for k, v in (payload.get("storyInput") or {}).items() if v}}
    payload["meta"].update(dashboard=DASHBOARD_FILE, today=datetime.now().strftime("%B %Y"))
    print("Data check")
    for k, v in report.items():
        print(f"  {k:<45} {v}")

    print("\n[3/3] Writing the pages ...")
    chart_tag = "<script>" + tpl["CHARTJS"].strip() + "\n</script>"
    dash, story = out_dir / DASHBOARD_FILE, out_dir / STORY_FILE
    write_page(tpl["DASHBOARD"], chart_tag, payload, dash)
    write_page(tpl["STORY"], chart_tag, payload, story)
    print(f"  {dash}\n  {story}\n\nDone.")
    if OPEN_WHEN_DONE and not args.no_open:
        webbrowser.open(story.as_uri())


if __name__ == "__main__":
    try:
        main()
    except PermissionError as e:
        print(f"\nA file is open in another program ({e.filename}). Close it (e.g. the workbook in Excel) and run again.")
    finally:
        if os.name == "nt" and "PROMPT" not in os.environ:   # double-clicked in Explorer: keep the window open
            input("\nPress Enter to close ...")
